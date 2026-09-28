from datetime import datetime, timedelta
from calendar import monthrange
import json
from .models import Client, Room, Reservation, Payment, Service
from django.shortcuts import render, redirect, get_object_or_404
from django.core.exceptions import ValidationError
from django.db.models import Q, Sum, F, Count, Avg, Max
from django.views.decorators.http import require_POST
from django.utils import timezone
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import Group
from .models import Room, Client, Reservation, Service, Payment, Expense, HotelSetting, Report, UserProfile, CalendarEvent
from .forms import ReservationForm, ClientForm, PaymentForm, ExpenseForm, HotelSettingForm, ServiceForm, RoomForm, ReportForm, HotelRegistrationForm, CalendarEventForm
from django.contrib.auth.models import User
from django.db import transaction
from django.urls import reverse
from django.utils import timezone
from django.db.models import Q, Sum


def _get_user_hotel(request):
    """Retourne l'hôtel de l'utilisateur connecté, ou None."""
    try:
        return request.user.profile.hotel
    except (UserProfile.DoesNotExist, AttributeError):
        return None


# Chemin menant à l'hôtel pour chaque modèle (cloisonnement des données entre hôtels)
HOTEL_LOOKUPS = {
    'Room': 'hotel',
    'Client': 'hotel',
    'Reservation': 'chambre__hotel',
    'Service': 'hotel',
    'Payment': 'reservation__chambre__hotel',
    'Expense': 'hotel',
    'CalendarEvent': 'hotel',
}


def _limiter_form_a_hotel(form, request):
    """Un formulaire ne doit proposer (ni accepter) que les chambres et clients de l'hôtel."""
    hotel = _get_user_hotel(request)
    if not hotel:
        return form
    if 'chambre' in form.fields:
        form.fields['chambre'].queryset = Room.objects.filter(hotel=hotel)
    if 'client' in form.fields:
        form.fields['client'].queryset = Client.objects.filter(hotel=hotel)
    if 'reservation' in form.fields:
        form.fields['reservation'].queryset = Reservation.objects.filter(chambre__hotel=hotel)
    return form


def _scoped(request, model_or_qs):
    """Limite une recherche aux données de l'hôtel de l'utilisateur.

    Un hôtel qui demande l'identifiant d'un autre hôtel obtient une page « introuvable ».
    Un compte sans hôtel (superutilisateur technique) garde l'accès complet.
    """
    qs = model_or_qs if hasattr(model_or_qs, 'model') else model_or_qs.objects.all()
    hotel = _get_user_hotel(request)
    if hotel is None:
        return qs
    return qs.filter(**{HOTEL_LOOKUPS[qs.model.__name__]: hotel})

# --- HELPERS PERMISSIONS ---
def _is_admin(user):
    """Vérifie si l'utilisateur appartient au groupe Admin."""
    return user.groups.filter(name='Admin').exists()

def _admin_required(view_func):
    """Décorateur : @login_required + vérification groupe Admin."""
    from functools import wraps
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return redirect('login')
        if not _is_admin(request.user):
            messages.error(request, "Accès refusé. Cette fonctionnalité est réservée à l'administrateur.")
            return redirect('home')
        return view_func(request, *args, **kwargs)
    return wrapper


def _can_manage_expenses(user):
    """Admin ou Gérant : accès à la gestion des dépenses."""
    return user.groups.filter(name__in=['Admin', 'Gerant']).exists()


def _expenses_access_required(view_func):
    """Décorateur : utilisateur authentifié + groupe Admin ou Gerant."""
    from functools import wraps
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return redirect('login')
        if not _can_manage_expenses(request.user):
            messages.error(
                request,
                "Accès refusé. La gestion des dépenses est réservée aux administrateurs et aux gérants."
            )
            return redirect('home')
        return view_func(request, *args, **kwargs)
    return wrapper

# --- AUTHENTICATION ---
def login_view(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        user = authenticate(request, username=username, password=password)
        if user is not None:
            login(request, user)
            return redirect('home')
        else:
            messages.error(request, 'Nom d\'utilisateur ou mot de passe incorrect.')
    response = render(request, 'core/login.html')
    # Efface l'ancien cookie « hôtel mémorisé » des navigateurs qui l'ont encore
    if 'hm_hotel' in request.COOKIES:
        response.delete_cookie('hm_hotel')
    return response

def register_view(request):
    """Enregistre un nouvel hôtel et son compte administrateur, puis renvoie vers la connexion."""
    if request.user.is_authenticated:
        return redirect('home')
    form = HotelRegistrationForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        data = form.cleaned_data
        with transaction.atomic():
            hotel = HotelSetting.objects.create(
                nom_etablissement=data['nom_etablissement'],
                adresse=data['adresse'],
                telephone=data['telephone'],
                email=data['email'],
                nom_gerant=data['nom_gerant'],
            )
            user = User.objects.create_user(
                username=data['username'],
                email=data['email'],
                password=data['password1'],
            )
            user.groups.add(Group.objects.get_or_create(name='Admin')[0])
            UserProfile.objects.create(user=user, hotel=hotel)
        messages.success(
            request,
            f"L'hôtel « {hotel.nom_etablissement} » a été enregistré. Connectez-vous avec vos identifiants."
        )
        return redirect('login')
    return render(request, 'core/register.html', {'form': form})

@login_required
def logout_view(request):
    logout(request)
    return redirect('login')

# --- DASHBOARD ---
def _room_nights_in_month(reservation, year, month):
    """Nombre de nuits d'une résa dans le mois donné."""
    start = datetime(year, month, 1).date()
    _, last_day = monthrange(year, month)
    end = datetime(year, month, last_day).date()
    if reservation.date_depart < start or reservation.date_arrivee > end:
        return 0
    return (min(reservation.date_depart, end) - max(reservation.date_arrivee, start)).days + 1

@login_required
def home(request):
    """
    Tableau de bord unique : le template distingue admin (finance complète)
    et gérant (vue opérationnelle + encaissements), via `is_admin` (context processor).
    """
    today = timezone.now().date()
    hotel = _get_user_hotel(request)

    # KPI Chambres
    rooms_qs = Room.objects.filter(hotel=hotel) if hotel else Room.objects.all()
    total_rooms = rooms_qs.count()
    occupees = rooms_qs.filter(statut_actuel='Occupée').count()
    disponibles = rooms_qs.filter(statut_actuel='Libre').count()
    en_maintenance = rooms_qs.filter(statut_actuel='En maintenance').count()
    taux_occupation = (occupees / total_rooms * 100) if total_rooms > 0 else 0

    # Revenus aujourd'hui (paiements enregistrés aujourd'hui)
    pay_qs = Payment.objects.filter(reservation__chambre__hotel=hotel) if hotel else Payment.objects.all()
    revenus_aujourdhui = pay_qs.filter(
        date_paiement__date=today
    ).aggregate(total=Sum('montant_paye'))['total'] or 0

    # Revenus ce mois
    revenus_mois = pay_qs.filter(
        date_paiement__year=today.year,
        date_paiement__month=today.month
    ).aggregate(total=Sum('montant_paye'))['total'] or 0

    # Dépenses ce mois
    exp_qs = Expense.objects.filter(hotel=hotel) if hotel else Expense.objects.all()
    depenses_mois = exp_qs.filter(
        date_depense__year=today.year,
        date_depense__month=today.month
    ).aggregate(total=Sum('montant'))['total'] or 0

    # Arrivées et départs aujourd'hui
    res_qs = Reservation.objects.filter(chambre__hotel=hotel) if hotel else Reservation.objects.all()
    arrivees_jour = res_qs.filter(date_arrivee=today, statut_reser='Confirmée')
    nb_arrivees = arrivees_jour.count()
    departs_jour = res_qs.filter(date_depart=today, statut_reser='En cours')
    nb_departs = departs_jour.count()

    # Clients et réservations totaux
    clients_qs = Client.objects.filter(hotel=hotel) if hotel else Client.objects.all()
    total_clients = clients_qs.count()
    reservations_en_cours = res_qs.filter(statut_reser='En cours').count()
    reservations_confirmees = res_qs.filter(statut_reser='Confirmée').count()

    # Revenus par mois (6 derniers mois)
    mois_labels = []
    revenus_mensuels = []
    depenses_mensuels = []
    taux_mensuels = []
    noms_mois = ['janv.', 'févr.', 'mars', 'avr.', 'mai', 'juin', 'juil.', 'août', 'sept.', 'oct.', 'nov.', 'déc.']
    for i in range(5, -1, -1):
        m = today.month - 1 - i
        year = today.year + (m // 12)
        month = (m % 12) + 1
        mois_labels.append(noms_mois[month - 1])
        rev = pay_qs.filter(
            date_paiement__year=year,
            date_paiement__month=month
        ).aggregate(total=Sum('montant_paye'))['total'] or 0
        revenus_mensuels.append(float(rev))
        dep = exp_qs.filter(
            date_depense__year=year,
            date_depense__month=month
        ).aggregate(total=Sum('montant'))['total'] or 0
        depenses_mensuels.append(float(dep))
        # Taux d'occupation du mois
        _, ndays = monthrange(year, month)
        capacity = total_rooms * ndays
        if capacity == 0:
            taux_mensuels.append(0)
        else:
            room_nights = 0
            for res in res_qs.filter(
                statut_reser__in=['Confirmée', 'En cours', 'Terminée'],
                date_arrivee__lte=datetime(year, month, ndays).date(),
                date_depart__gte=datetime(year, month, 1).date(),
            ):
                room_nights += _room_nights_in_month(res, year, month)
            taux_mensuels.append(round(room_nights / capacity * 100, 1))

    # Répartition des chambres par type
    repartition_chambres = list(
        rooms_qs.values('type_chambre').annotate(nb=Count('id')).order_by('type_chambre')
    )

    # Répartition des chambres par statut
    repartition_statuts = list(
        rooms_qs.values('statut_actuel').annotate(nb=Count('id')).order_by('statut_actuel')
    )

    # Dernières réservations
    dernieres_reservations = res_qs.select_related('client', 'chambre').order_by('-date_reservation')[:5]

    # Dernières dépenses
    dernieres_depenses = exp_qs.order_by('-date_depense')[:5]

    context = {
        'taux_occupation': taux_occupation,
        'disponibles': disponibles,
        'occupees': occupees,
        'en_maintenance': en_maintenance,
        'total_rooms': total_rooms,
        'revenus_aujourdhui': int(revenus_aujourdhui),
        'revenus_mois': int(revenus_mois),
        'depenses_mois': int(depenses_mois),
        'benefice_mois': int(revenus_mois) - int(depenses_mois),
        'nb_arrivees': nb_arrivees,
        'nb_departs': nb_departs,
        'total_clients': total_clients,
        'reservations_en_cours': reservations_en_cours,
        'reservations_confirmees': reservations_confirmees,
        'arrivees': arrivees_jour,
        'departs': departs_jour,
        'mois_labels': json.dumps(mois_labels),
        'revenus_mensuels': json.dumps(revenus_mensuels),
        'depenses_mensuels': json.dumps(depenses_mensuels),
        'taux_mensuels': json.dumps(taux_mensuels),
        'repartition_chambres': repartition_chambres,
        'repartition_statuts': repartition_statuts,
        'dernieres_reservations': dernieres_reservations,
        'dernieres_depenses': dernieres_depenses,
    }
    return render(request, 'core/dashboard.html', context)

# --- CHAMBRES ---
@login_required
def rooms_list(request):
    hotel = _get_user_hotel(request)
    q = request.GET.get('q', '')
    type_filter = request.GET.get('type')
    status_filter = request.GET.get('status')
    rooms = Room.objects.filter(hotel=hotel) if hotel else Room.objects.all()
    if q:
        rooms = rooms.filter(numero_chambre__icontains=q)
    if type_filter:
        rooms = rooms.filter(type_chambre=type_filter)
    if status_filter:
        rooms = rooms.filter(statut_actuel=status_filter)
    rooms = rooms.order_by('numero_chambre')
    
    context = {'rooms': rooms, 'q': q, 'type_filter': type_filter, 'status_filter': status_filter}
    return render(request, 'core/chambres/list.html', context)

@_admin_required
def room_create(request):
    hotel = _get_user_hotel(request)
    if request.method == 'POST':
        form = RoomForm(request.POST)
        if form.is_valid():
            room = form.save(commit=False)
            room.hotel = hotel
            room.save()
            messages.success(request, "Chambre créée avec succès.")
            return redirect('rooms_list')
        else:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f"{error}")
    else:
        form = RoomForm()
    
    return render(request, 'core/chambres/form.html', {'form': form})

@login_required
def room_detail(request, pk):
    room = get_object_or_404(_scoped(request, Room), pk=pk)
    return render(request, 'core/chambres/detail.html', {'room': room})

@_admin_required
def room_update(request, pk):
    room = get_object_or_404(_scoped(request, Room), pk=pk)
    if request.method == 'POST':
        form = RoomForm(request.POST, instance=room)
        if form.is_valid():
            form.save()
            messages.success(request, "Chambre mise à jour avec succès.")
            return redirect('rooms_list')
        else:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f"{error}")
    else:
        form = RoomForm(instance=room)
    
    return render(request, 'core/chambres/form.html', {'form': form, 'room': room})

@_admin_required
def room_delete(request, pk):
    room = get_object_or_404(_scoped(request, Room), pk=pk)
    if request.method == 'POST':
        numero = room.numero_chambre
        room.delete()
        messages.warning(request, f"Chambre {numero} supprimée avec succès.")
        return redirect('rooms_list')
    return render(request, 'core/chambres/confirm_delete.html', {'room': room})


# --- CLIENTS ---
@login_required
def clients_list(request):
    from django.utils import timezone
    from django.db.models import Q

    hotel = _get_user_hotel(request)
    today = timezone.now().date()
    first_day = today.replace(day=1)

    q = request.GET.get('q', '').strip()
    clients_qs = Client.objects.filter(hotel=hotel) if hotel else Client.objects.all()
    clients = clients_qs
    if q:
        filtre = (
            Q(nom__icontains=q) | Q(prenom__icontains=q) | Q(email__icontains=q)
            | Q(telephone__icontains=q) | Q(numero_document__icontains=q)
        )
        # Recherche par numéro client : "#00003", "00003" ou "3"
        numero = q.lstrip('#')
        if numero.isdigit():
            filtre |= Q(numero_client=int(numero))
        clients = clients.filter(filtre)
    # Client régulier (comme la démo) : au moins 2 réservations
    clients = clients.annotate(nb_reservations=Count('reservations', distinct=True)).order_by('numero_client')
    filtre_clients = request.GET.get('filtre', 'tous')
    if filtre_clients == 'reguliers':
        clients = clients.filter(nb_reservations__gte=2)

    stats = {
        'total':              clients_qs.count(),
        'arrives_aujourdhui': clients_qs.filter(date_entree=today).count(),
        'actifs':             clients_qs.filter(date_entree__lte=today, date_sortie__gte=today).count(),
        'reguliers':          clients_qs.annotate(nb=Count('reservations', distinct=True)).filter(nb__gte=2).count(),
        'nouveaux_mois':      clients_qs.filter(date_creation__date__gte=first_day).count(),
    }

    # Numéro que recevra le prochain client (affiché dans « Ajouter un client », comme la démo)
    dernier = clients_qs.aggregate(m=Max('numero_client'))['m'] or 0
    return render(request, 'core/clients/list.html', {
        'clients': clients, 'q': q, 'stats': stats, 'filtre': filtre_clients,
        'prochain_code': f"#{dernier + 1:05d}",
    })

@_admin_required
def client_create(request):
    hotel = _get_user_hotel(request)
    if request.method == 'POST':
        form = ClientForm(request.POST)
        if form.is_valid():
            client = form.save(commit=False)
            client.hotel = hotel
            client.save()
            messages.success(request, f"Client {client.code_client} créé avec succès.")
            return redirect('clients_list')
        else:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f"{error}")
    else:
        form = ClientForm()
    
    return render(request, 'core/clients/client_form.html', {'form': form})

@login_required
def client_detail(request, pk):
    client = get_object_or_404(_scoped(request, Client), pk=pk)
    return render(request, 'core/clients/detail.html', {'client': client})

@_admin_required
def client_update(request, pk):
    client = get_object_or_404(_scoped(request, Client), pk=pk)
    if request.method == 'POST':
        form = ClientForm(request.POST, instance=client)
        if form.is_valid():
            form.save()
            messages.success(request, "Client mis à jour avec succès.")
            return redirect('clients_list')
        else:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f"{error}")
    else:
        form = ClientForm(instance=client)
    
    return render(request, 'core/clients/client_form.html', {'form': form, 'client': client})

@_admin_required
def client_delete(request, pk):
    client = get_object_or_404(_scoped(request, Client), pk=pk)
    if request.method == 'POST':
        client.delete()
        messages.warning(request, f"Client {client.prenom} {client.nom} supprimé avec succès.")
        return redirect('clients_list')
    return render(request, 'core/clients/confirm_delete.html', {'client': client})

@login_required
def client_pdf(request, pk):
    from django.template.loader import render_to_string
    from django.http import HttpResponse
    from xhtml2pdf import pisa
    from io import BytesIO

    client = get_object_or_404(_scoped(request, Client), pk=pk)
    html_string = render_to_string('core/clients/fiche_pdf.html', {'client': client})

    # Créer le PDF
    result = BytesIO()
    pdf = pisa.pisaDocument(BytesIO(html_string.encode("UTF-8")), result)
    
    if not pdf.err:
        response = HttpResponse(result.getvalue(), content_type='application/pdf')
        # ?apercu=1 : affichage dans la fenêtre « Fiche client » ; sinon téléchargement
        mode = 'inline' if request.GET.get('apercu') else 'attachment'
        response['Content-Disposition'] = f'{mode}; filename="fiche_client_{client.nom}_{client.prenom}.pdf"'
        return response
    
    return HttpResponse('Erreur lors de la génération du PDF', status=500)


# --- RESERVATIONS ---
def _services_catalog(hotel):
    """Services proposés à la création d'une réservation : uniquement ceux créés
    par l'hôtel dans la page Services (sans réservation), un par nom. Liste vide sinon."""
    qs = Service.objects.filter(reservation__isnull=True)
    if hotel:
        qs = qs.filter(hotel=hotel)
    catalog = {}
    for s in qs.order_by('-id'):
        catalog.setdefault(s.nom_service.strip().lower(), {'nom': s.nom_service.strip(), 'prix': s.prix_service})
    return sorted(catalog.values(), key=lambda s: s['nom'].lower())


@login_required
def reservations_list(request):
    hotel = _get_user_hotel(request)
    status = request.GET.get('status')
    reservations = Reservation.objects.select_related('client', 'chambre').filter(chambre__hotel=hotel) if hotel else Reservation.objects.select_related('client', 'chambre').all()
    if status:
        reservations = reservations.filter(statut_reser=status)
    clients = Client.objects.filter(hotel=hotel) if hotel else Client.objects.all()
    chambres_libres = Room.objects.filter(statut_actuel='Libre', hotel=hotel) if hotel else Room.objects.filter(statut_actuel='Libre')
    return render(request, 'core/reservations/list.html', {
        'reservations': reservations, 'status': status,
        'clients': clients, 'chambres_libres': chambres_libres,
        'services_catalog': _services_catalog(hotel),
    })

@login_required
def reservation_create(request):
    hotel = _get_user_hotel(request)
    form = _limiter_form_a_hotel(ReservationForm(request.POST or None), request)

    if request.method == 'POST':
        if form.is_valid():
            res = form.save(commit=False)

            # 1. Gestion du NOUVEAU client
            if form.cleaned_data.get('nouveau_client'):
                client = Client(
                    hotel=hotel,
                    nom=form.cleaned_data.get('nouveau_nom'),
                    prenom=form.cleaned_data.get('nouveau_prenom'),
                    email=form.cleaned_data.get('nouveau_email') or None,
                    telephone=form.cleaned_data.get('nouveau_telephone') or '',
                    adresse=form.cleaned_data.get('nouveau_adresse') or '',
                    type_document=form.cleaned_data.get('nouveau_type_document'),
                    numero_document=form.cleaned_data.get('nouveau_numero_document'),
                    # FIX : On passe les dates de la résa au client car elles sont obligatoires dans ton modèle Client
                    date_entree=form.cleaned_data.get('date_arrivee'),
                    date_sortie=form.cleaned_data.get('date_depart'),
                )
                try:
                    client.full_clean()
                    client.save()
                    res.client = client 
                except ValidationError as e:
                    for msg in e.messages:
                        messages.error(request, f"Erreur client : {msg}")
                    return render(request, 'core/reservations/form.html', {'form': form})
            
            # 2. Gestion du client EXISTANT (si pas de nouveau client coché)
            else:
                res.client = form.cleaned_data.get('client')

            # 3. Logique prix et sauvegarde
            if res.chambre:
                res.prix_nuit = res.chambre.prix_nuit
            
            res.montant_avance = 0

            try:
                res.full_clean()
                with transaction.atomic():
                    res.save()
                    # Services supplémentaires cochés : prix repris du catalogue, pas du navigateur
                    catalog = {s['nom'].lower(): s for s in _services_catalog(hotel)}
                    for nom in dict.fromkeys(request.POST.getlist('services')):
                        s = catalog.get(nom.strip().lower())
                        if s:
                            Service.objects.create(reservation=res, nom_service=s['nom'], prix_service=s['prix'])
                messages.success(request, "Réservation effectuée.")
                return redirect('reservation_detail', pk=res.pk)
            except ValidationError as e:
                # Si erreur ici, c'est souvent un chevauchement de dates (ton clean() dans models.py)
                for field, errors in e.message_dict.items():
                    messages.error(request, f"{field}: {', '.join(errors)}")
        else:
            # Affiche les erreurs de formulaire (ex: client manquant si nouveau_client non coché)
            for field, errors in form.errors.items():
                messages.error(request, f"Erreur dans {field}: {errors.as_text()}")

    return render(request, 'core/reservations/form.html', {'form': form})



@login_required
def reservation_detail(request, pk):
    res = get_object_or_404(_scoped(request, Reservation.objects.select_related('client', 'chambre')), pk=pk)
    # On récupère les services et paiements liés
    services = res.services.all()
    paiements = res.paiements.all()
    return render(request, 'core/reservations/detail.html', {
        'r': res, 
        'services': services,
        'paiements': paiements,
        'service_form': ServiceForm(), 
        'payment_form': PaymentForm()
    })

@login_required
def reservation_update(request, pk):
    res = get_object_or_404(_scoped(request, Reservation), pk=pk)
    form = _limiter_form_a_hotel(ReservationForm(request.POST or None, instance=res), request)
    if form.is_valid():
        res = form.save(commit=False)
        # Met à jour le prix_nuit si la chambre a changé
        if res.chambre and hasattr(res.chambre, 'prix_nuit'):
            res.prix_nuit = res.chambre.prix_nuit
        res.montant_avance = 0
        res.save()
        return redirect('reservation_detail', pk=pk)
    return render(request, 'core/reservations/form.html', {'form': form, 'edit_mode': True})

@login_required
def reservation_delete(request, pk):
    res = get_object_or_404(_scoped(request, Reservation), pk=pk)
    if request.method == 'POST':
        res.delete()
        messages.warning(request, "Réservation supprimée.")
        return redirect('reservations_list')
    return render(request, 'core/reservations/confirm_delete.html', {'reservation': res})

# --- ACTIONS CHECK-IN / CHECK-OUT ---
@login_required
@require_POST
def reservation_checkin(request, pk):
    res = get_object_or_404(_scoped(request, Reservation), pk=pk)
    res.statut_reser = 'En cours' # Aligné avec votre STATUS_CHOICES
    res.save()
    messages.info(request, "Client enregistré (Check-in).")
    return redirect('reservation_detail', pk=pk)

@login_required
@require_POST
def reservation_checkout(request, pk):
    res = get_object_or_404(_scoped(request, Reservation.objects.select_related('client', 'chambre')), pk=pk)
    if res.solde_du > 0:
        messages.error(
            request,
            f"Check-out impossible : solde restant de {res.solde_du:,.0f} FCFA. "
            f"Le client doit régler la totalité avant de quitter."
        )
        return redirect('reservation_detail', pk=pk)
    res.statut_reser = 'Terminée'
    res.save()
    messages.info(request, "Client sorti (Check-out).")
    return redirect('reservation_detail', pk=pk)

@login_required
@require_POST
def reservation_cancel(request, pk):
    res = get_object_or_404(_scoped(request, Reservation), pk=pk)
    res.delete()
    messages.warning(request, "Réservation annulée et supprimée.")
    return redirect('reservations_list')

def _fcfa(value):
    """12500 -> '12 500' (espace simple, lisible dans les PDF)."""
    return f"{int(round(float(value or 0))):,}".replace(",", " ")


def _render_pdf(template, context, filename, inline=True):
    from django.template.loader import render_to_string
    from django.http import HttpResponse
    from xhtml2pdf import pisa
    from io import BytesIO

    result = BytesIO()
    pdf = pisa.pisaDocument(BytesIO(render_to_string(template, context).encode("UTF-8")), result)
    if pdf.err:
        return HttpResponse("Erreur lors de la génération du PDF", status=500)
    response = HttpResponse(result.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'{"inline" if inline else "attachment"}; filename="{filename}"'
    return response


@login_required
def reservation_invoice(request, pk):
    r = get_object_or_404(_scoped(request, Reservation.objects.select_related('client', 'chambre', 'chambre__hotel')), pk=pk)

    nb_nuits = r.nb_nuits
    prix_nuit = float(r.prix_nuit or 0)
    total_chambre = nb_nuits * prix_nuit
    total_services = float(r.total_services or 0)
    total_general = total_chambre + total_services
    total_paye = float(r.total_paye or 0)
    solde_du = float(r.solde_du or 0)

    lignes = [{
        'libelle': f"Hébergement — Chambre {r.chambre.numero_chambre} ({r.chambre.type_chambre})",
        'detail': f"du {r.date_arrivee:%d/%m/%Y} au {r.date_depart:%d/%m/%Y}",
        'qte': nb_nuits, 'pu': _fcfa(prix_nuit), 'total': _fcfa(total_chambre),
    }] + [{
        'libelle': s.nom_service, 'detail': '', 'qte': 1,
        'pu': _fcfa(s.prix_service), 'total': _fcfa(s.prix_service),
    } for s in r.services.all()]

    paiements = [{
        'date': p.date_paiement, 'methode': p.methode_paiement,
        'reference': p.reference_paiement, 'montant': _fcfa(p.montant_paye),
    } for p in r.paiements.order_by('date_paiement')]

    return _render_pdf('core/reservations/facture_pdf.html', {
        'r': r,
        'hotel': r.chambre.hotel,
        'numero': f"FAC-{r.date_reservation:%Y}-{r.id:04d}",
        'lignes': lignes,
        'paiements': paiements,
        'nb_nuits': nb_nuits,
        'total_chambre': _fcfa(total_chambre),
        'total_services': _fcfa(total_services),
        'total_general': _fcfa(total_general),
        'total_paye': _fcfa(total_paye),
        'solde_du': _fcfa(max(solde_du, 0)),
        'est_solde': solde_du <= 0,
    }, f"facture_reservation_{r.id}.pdf")


@login_required
def payment_receipt(request, pk):
    """Reçu de paiement (PDF) : un par encaissement."""
    hotel = _get_user_hotel(request)
    qs = Payment.objects.select_related('reservation', 'reservation__client', 'reservation__chambre', 'reservation__chambre__hotel')
    if hotel:
        qs = qs.filter(reservation__chambre__hotel=hotel)
    p = get_object_or_404(qs, pk=pk)
    r = p.reservation
    solde_du = float(r.solde_du or 0)
    return _render_pdf('core/payments/recu_pdf.html', {
        'p': p,
        'r': r,
        'hotel': r.chambre.hotel,
        'numero': f"REC-{p.date_paiement:%Y}-{p.id:04d}",
        'montant': _fcfa(p.montant_paye),
        'total_facture': _fcfa(float(r.total_chambre or 0) + float(r.total_services or 0)),
        'total_paye': _fcfa(r.total_paye),
        'solde_du': _fcfa(max(solde_du, 0)),
        'est_solde': solde_du <= 0,
    }, f"recu_paiement_{p.id}.pdf")

# --- SERVICES ET PAIEMENTS ---
@login_required
@require_POST
def reservation_add_service(request, pk):
    res = get_object_or_404(_scoped(request, Reservation), pk=pk)
    form = ServiceForm(request.POST)
    if form.is_valid():
        s = form.save(commit=False)
        s.reservation = res
        s.save()
    return redirect('reservation_detail', pk=pk)

@login_required
@require_POST
def reservation_delete_service(request, pk, service_pk):
    """Retire un service de la réservation (depuis la fiche de réservation)."""
    hotel = _get_user_hotel(request)
    qs = Service.objects.filter(reservation_id=pk)
    if hotel:
        qs = qs.filter(reservation__chambre__hotel=hotel)
    service = get_object_or_404(qs, pk=service_pk)
    nom = service.nom_service
    service.delete()
    messages.success(request, f"Service « {nom} » supprimé.")
    return redirect('reservation_detail', pk=pk)

@login_required
@require_POST
def reservation_add_payment(request, pk):
    res = get_object_or_404(_scoped(request, Reservation), pk=pk)
    # On initialise le formulaire avec une instance liée à la réservation
    # pour que la validation clean() puisse accéder à res.solde_du
    form = PaymentForm(request.POST, instance=Payment(reservation=res))
    
    if form.is_valid():
        p = form.save(commit=False)
        p.reservation = res
        try:
            p.full_clean()
            p.save()
            messages.success(request, f"Paiement de {p.montant_paye} FCFA enregistré.")
        except ValidationError as e:
            for msg in e.messages:
                messages.error(request, msg)
    else:
        for field, errors in form.errors.items():
            for error in errors:
                messages.error(request, f"{error}")
    return redirect('reservation_detail', pk=pk)

# --- DEPENSES ---
@_expenses_access_required
def expenses_list(request):
    hotel = _get_user_hotel(request)
    q = request.GET.get('q', '')
    cat = request.GET.get('cat', '')
    expenses = Expense.objects.filter(hotel=hotel) if hotel else Expense.objects.all()
    if q:
        expenses = expenses.filter(Q(description__icontains=q))
    if cat:
        expenses = expenses.filter(categorie=cat)
    return render(request, 'core/expenses/list.html', {'expenses': expenses, 'q': q, 'cat': cat, 'categories': Expense.CATEGORIES})

@_expenses_access_required
def expenses_create(request):
    hotel = _get_user_hotel(request)
    form = ExpenseForm(request.POST or None, request.FILES or None)
    if request.method == 'POST':
        if form.is_valid():
            expense = form.save(commit=False)
            expense.hotel = hotel
            expense.save()
            messages.success(request, "Dépense enregistrée.")
        else:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f"{error}")
        return redirect('expenses_list')
    return render(request, 'core/expenses/form.html', {'form': form})

# --- CALENDRIER ---
MOIS_FR = ['janvier', 'février', 'mars', 'avril', 'mai', 'juin', 'juillet',
           'août', 'septembre', 'octobre', 'novembre', 'décembre']


@login_required
def calendar_view(request):
    """Calendrier mensuel des événements (même fonctionnement que la démo)."""
    import calendar as calendar_module

    hotel = _get_user_hotel(request)
    today = timezone.localdate()
    try:
        year, month = (int(x) for x in request.GET.get('mois', '').split('-'))
        first_day = datetime(year, month, 1).date()
    except (ValueError, TypeError):
        year, month = today.year, today.month
        first_day = today.replace(day=1)
    last_day = first_day.replace(day=monthrange(year, month)[1])

    events_qs = CalendarEvent.objects.select_related('client', 'chambre').filter(date__range=(first_day, last_day))
    if hotel:
        events_qs = events_qs.filter(hotel=hotel)
    events = list(events_qs)

    # Réservations du mois (hors annulées), affichées sur chaque nuit du séjour
    resa_qs = Reservation.objects.select_related('client', 'chambre').filter(
        date_arrivee__lte=last_day, date_depart__gte=first_day
    ).exclude(statut_reser='Annulée').order_by('date_arrivee', 'chambre__numero_chambre')
    if hotel:
        resa_qs = resa_qs.filter(chambre__hotel=hotel)
    reservations = list(resa_qs)
    resa_css = {'Confirmée': 'confirmee', 'En cours': 'en_cours', 'Terminée': 'terminee'}

    def _day_resas(day):
        return [r for r in reservations
                if r.date_arrivee <= day < r.date_depart or r.date_arrivee == r.date_depart == day]

    weeks = []
    for week in calendar_module.Calendar(firstweekday=6).monthdatescalendar(year, month):  # du dimanche au samedi
        weeks.append([])
        for day in week:
            day_resas = _day_resas(day)
            day_events = [e for e in events if e.date == day]
            items = [{
                'kind': 'resa', 'obj': r, 'css': resa_css.get(r.statut_reser, 'confirmee'),
                'is_arrival': r.date_arrivee == day,
            } for r in day_resas] + [{'kind': 'event', 'obj': e, 'css': e.statut} for e in day_events]
            weeks[-1].append({
                'date': day,
                'in_month': day.month == month,
                'is_today': day == today,
                'events': day_events,
                'items': items,
                'more': max(len(items) - 3, 0),
            })

    prev_month = (first_day - timedelta(days=1)).replace(day=1)
    next_month = last_day + timedelta(days=1)
    context = {
        'weeks': weeks,
        'events': events,
        'nb_reservations': len(reservations),
        'month_label': f"{MOIS_FR[month - 1]} {year}",
        'month_param': first_day.strftime('%Y-%m'),
        'prev_month': prev_month.strftime('%Y-%m'),
        'next_month': next_month.strftime('%Y-%m'),
        'weekdays': ['Dim', 'Lun', 'Mar', 'Mer', 'Jeu', 'Ven', 'Sam'],
        'statuts': CalendarEvent.STATUS_CHOICES,
        'clients': Client.objects.filter(hotel=hotel) if hotel else Client.objects.all(),
        'chambres': Room.objects.filter(hotel=hotel) if hotel else Room.objects.all(),
    }
    return render(request, 'core/calendar.html', context)


def _calendar_redirect(request, day=None):
    """Retour au calendrier, sur le mois de l'événement (ou celui d'où vient le formulaire)."""
    mois = day.strftime('%Y-%m') if day else request.POST.get('mois', '')
    url = reverse('calendar')
    return redirect(f"{url}?mois={mois}" if mois else url)


@login_required
@require_POST
def calendar_event_create(request):
    hotel = _get_user_hotel(request)
    form = CalendarEventForm(request.POST, hotel=hotel)
    if form.is_valid():
        event = form.save(commit=False)
        event.hotel = hotel
        event.save()
        messages.success(request, "Événement ajouté au calendrier.")
        return _calendar_redirect(request, event.date)
    for errors in form.errors.values():
        for error in errors:
            messages.error(request, error)
    return _calendar_redirect(request)


def _get_event(request, pk):
    hotel = _get_user_hotel(request)
    if hotel:
        return get_object_or_404(CalendarEvent, pk=pk, hotel=hotel)
    return get_object_or_404(CalendarEvent, pk=pk)


@login_required
@require_POST
def calendar_event_status(request, pk):
    event = _get_event(request, pk)
    statut = request.POST.get('statut')
    if statut in dict(CalendarEvent.STATUS_CHOICES):
        event.statut = statut
        event.save(update_fields=['statut'])
        messages.success(request, f"Statut mis à jour : {event.get_statut_display()}.")
    return _calendar_redirect(request, event.date)


@login_required
@require_POST
def calendar_event_delete(request, pk):
    event = _get_event(request, pk)
    day = event.date
    event.delete()
    messages.warning(request, "Événement supprimé.")
    return _calendar_redirect(request, day)

# --- RAPPORTS ---
@_admin_required
def reports(request):
    date_debut = request.GET.get('date_debut', '')
    date_fin = request.GET.get('date_fin', '')
    generated = False

    base_context = {
        'taux_occupation': 0,
        'total_recettes': 0,
        'total_depenses': 0,
        'benefice_net': 0,
        'nb_reservations': 0,
        'nb_clients': 0,
        'nb_paiements': 0,
        'montant_moyen_paiement': 0,
        'revenu_moyen_par_reservation': 0,
        'revenu_moyen_par_client': 0,
        'depense_moyenne_journaliere': 0,
        'marge_beneficiaire': 0,
        'room_nights': 0,
        'capacity': 0,
        'revpar': 0,
        'encours_impayes': 0,
        'depenses_par_categorie': [],
        'reservations_par_statut': [],
        'paiements_par_methode': [],
        'top_clients': [],
        'top_chambres': [],
        'top_services': [],
        'daily_finance': [],
    }

    if date_debut and date_fin:
        try:
            data = _get_report_data(request)
            if data:
                generated = True
                base_context.update(data)

        except (ValueError, TypeError):
            messages.error(request, "Format de date invalide.")

    context = {
        **base_context,
        'date_debut': date_debut,
        'date_fin': date_fin,
        'generated': generated,
        'rapports': Report.objects.all().order_by('-date_debut'),
        'report_form': ReportForm(),
    }
    return render(request, 'core/reports/index.html', context)


@_admin_required
@require_POST
def report_upload(request):
    form = ReportForm(request.POST, request.FILES)
    if form.is_valid():
        rapport = form.save(commit=False)
        rapport.hotel = _get_user_hotel(request)
        rapport.save()
        messages.success(request, "Rapport uploadé avec succès.")
    else:
        for field, errors in form.errors.items():
            for error in errors:
                messages.error(request, f"{error}")
    return redirect('reports')

# --- PARAMETRES ---
@_admin_required
def settings_view(request):
    hotel = _get_user_hotel(request)
    if hotel:
        obj = hotel
    else:
        obj, created = HotelSetting.objects.get_or_create(pk=1)
    if request.method == 'POST':
        form = HotelSettingForm(request.POST, request.FILES, instance=obj)
        if form.is_valid():
            form.save()
            messages.success(request, "Paramètres mis à jour.")
            return redirect('settings_view')
    else:
        form = HotelSettingForm(instance=obj)
    return render(request, 'core/settings/index.html', {'form': form})

@_admin_required
@require_POST
def settings_sample_data(request):
    """Ajoute des données d'exemple (chambres, clients, réservations…) à l'hôtel de l'utilisateur."""
    from .sample_data import add_sample_data
    hotel = _get_user_hotel(request) or HotelSetting.objects.get_or_create(pk=1)[0]
    if add_sample_data(hotel):
        messages.success(request, "Données d'exemple ajoutées ! Consultez les différentes sections.")
    else:
        messages.info(request, "Les données d'exemple ont déjà été ajoutées à cet hôtel.")
    return redirect('settings_view')

# --- ROOMS UPDATE / DELETE ---
@_admin_required
def room_update(request, pk):
    room = get_object_or_404(_scoped(request, Room), pk=pk)
    form = RoomForm(request.POST or None, instance=room)
    if form.is_valid():
        form.save()
        messages.success(request, "Chambre mise à jour.")
        return redirect('rooms_list')
    return render(request, 'core/chambres/form.html', {'form': form, 'room': room})

@_admin_required
def room_delete(request, pk):
    room = get_object_or_404(_scoped(request, Room), pk=pk)
    if request.method == 'POST':
        room.delete()
        messages.warning(request, "Chambre supprimée.")
        return redirect('rooms_list')
    return render(request, 'core/chambres/confirm_delete.html', {'room': room})

# --- CLIENTS UPDATE / DELETE ---
@_admin_required
def client_update(request, pk):
    client = get_object_or_404(_scoped(request, Client), pk=pk)
    form = ClientForm(request.POST or None, instance=client)
    if form.is_valid():
        form.save()
        messages.success(request, "Client mis à jour.")
        return redirect('clients_list')
    return render(request, 'core/clients/client_form.html', {'form': form, 'client': client})

@_admin_required
def client_delete(request, pk):
    client = get_object_or_404(_scoped(request, Client), pk=pk)
    if request.method == 'POST':
        client.delete()
        messages.warning(request, "Client supprimé.")
        return redirect('clients_list')
    return render(request, 'core/clients/confirm_delete.html', {'client': client})

# --- SERVICES ---
@login_required
def services_list(request):
    hotel = _get_user_hotel(request)
    q = request.GET.get('q', '')
    services = Service.objects.filter(hotel=hotel) if hotel else Service.objects.all()
    if q:
        services = services.filter(Q(nom_service__icontains=q) | Q(description__icontains=q))
    return render(request, 'core/services/list.html', {'services': services, 'q': q})

@_admin_required
def service_create(request):
    if request.method == 'POST':
        form = ServiceForm(request.POST)
        if form.is_valid():
            service = form.save(commit=False)
            service.hotel = _get_user_hotel(request)
            service.save()
            messages.success(request, "Service créé avec succès.")
            return redirect('services_list')
    else:
        form = ServiceForm()
    return render(request, 'core/services/form.html', {'form': form})

@login_required
def service_detail(request, pk):
    service = get_object_or_404(_scoped(request, Service), pk=pk)
    return render(request, 'core/services/detail.html', {'service': service})

@_admin_required
def service_update(request, pk):
    service = get_object_or_404(_scoped(request, Service), pk=pk)
    form = ServiceForm(request.POST or None, instance=service)
    if form.is_valid():
        form.save()
        messages.success(request, "Service mis à jour.")
        return redirect('services_list')
    return render(request, 'core/services/form.html', {'form': form, 'service': service})

@_admin_required
def service_delete(request, pk):
    service = get_object_or_404(_scoped(request, Service), pk=pk)
    if request.method == 'POST':
        service.delete()
        messages.warning(request, "Service supprimé.")
        return redirect('services_list')
    return render(request, 'core/services/confirm_delete.html', {'service': service})

# --- EXPENSES UPDATE / DELETE ---
@_expenses_access_required
def expenses_update(request, pk):
    expense = get_object_or_404(_scoped(request, Expense), pk=pk)
    form = ExpenseForm(request.POST or None, request.FILES or None, instance=expense)
    if form.is_valid():
        form.save()
        messages.success(request, "Dépense mise à jour.")
        return redirect('expenses_list')
    return render(request, 'core/expenses/form.html', {'form': form, 'expense': expense})

@_expenses_access_required
def expenses_delete(request, pk):
    expense = get_object_or_404(_scoped(request, Expense), pk=pk)
    if request.method == 'POST':
        expense.delete()
        messages.warning(request, "Dépense supprimée.")
        return redirect('expenses_list')
    return render(request, 'core/expenses/confirm_delete.html', {'expense': expense})

# PAIEMENTS UPDATE / DELETE ---

@login_required
def payment_list(request):  
    today = timezone.now().date()
    
    hotel = _get_user_hotel(request)
    pay_base = Payment.objects.filter(reservation__chambre__hotel=hotel) if hotel else Payment.objects.all()
    payments = pay_base.select_related(
        'reservation', 
        'reservation__client',
        'reservation__chambre'
    ).order_by('-date_paiement')
    
    # Filtrage par recherche
    q = request.GET.get('q', '')
    if q:
        payments = payments.filter(
            Q(reservation__client__nom__icontains=q) |
            Q(reservation__client__prenom__icontains=q) |
            Q(reference_paiement__icontains=q)
        )
    
    # Calcul des stats
    total_encaisse = pay_base.aggregate(
        total=Sum('montant_paye')
    )['total'] or 0
    
    paiements_aujourdhui = pay_base.filter(
        date_paiement__date=today
    ).count()

    # Réservations actives pour le formulaire d'ajout
    res_base = Reservation.objects.filter(chambre__hotel=hotel) if hotel else Reservation.objects.all()
    reservations = res_base.select_related('client', 'chambre').filter(
        statut_reser__in=['Confirmée', 'En cours']
    ).order_by('-date_reservation')
    
    return render(request, 'core/payments/list.html', {
        'payments': payments,
        'q': q,
        'total_encaisse': total_encaisse,
        'paiements_aujourdhui': paiements_aujourdhui,
        'reservations': reservations,
        'payment_form': PaymentForm(),
    })

@login_required
def payment_create(request):
    if request.method == 'POST':
        reservation_id = request.POST.get('reservation')
        reservation = get_object_or_404(_scoped(request, Reservation), pk=reservation_id)
        form = _limiter_form_a_hotel(PaymentForm(request.POST), request)
        if form.is_valid():
            payment = form.save(commit=False)
            payment.reservation = reservation
            payment.save()
            messages.success(request, f"Paiement de {payment.montant_paye} FCFA enregistré avec succès.")
        else:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f"{error}")
    return redirect('payment_list')


# --- EXPORT RAPPORTS PDF / EXCEL ---
def _get_report_data(request):
    """Calcule les données de rapport à partir des paramètres GET."""
    date_debut = request.GET.get('date_debut', '')
    date_fin = request.GET.get('date_fin', '')
    if not date_debut or not date_fin:
        return None

    try:
        d_debut = datetime.strptime(date_debut, '%Y-%m-%d').date()
        d_fin = datetime.strptime(date_fin, '%Y-%m-%d').date()
    except (TypeError, ValueError):
        return None
    if d_fin < d_debut:
        return None

    hotel = _get_user_hotel(request)
    rooms_qs = Room.objects.filter(hotel=hotel) if hotel else Room.objects.all()
    total_rooms = rooms_qs.count()

    pay_qs = Payment.objects.filter(reservation__chambre__hotel=hotel) if hotel else Payment.objects.all()
    pay_period = pay_qs.filter(
        date_paiement__date__gte=d_debut,
        date_paiement__date__lte=d_fin
    )
    recettes = pay_period.aggregate(total=Sum('montant_paye'))['total'] or 0

    exp_qs = Expense.objects.filter(hotel=hotel) if hotel else Expense.objects.all()
    exp_period = exp_qs.filter(
        date_depense__gte=d_debut,
        date_depense__lte=d_fin
    )
    depenses_total = exp_period.aggregate(total=Sum('montant'))['total'] or 0

    benefice_net = recettes - depenses_total

    res_qs = Reservation.objects.filter(chambre__hotel=hotel) if hotel else Reservation.objects.all()
    res_period = res_qs.filter(
        date_arrivee__lte=d_fin,
        date_depart__gte=d_debut
    )
    nb_reservations = res_period.count()

    nb_clients = res_period.values('client').distinct().count()

    nb_jours = (d_fin - d_debut).days + 1
    capacity = total_rooms * nb_jours
    taux_occupation = 0
    room_nights = 0
    if capacity > 0:
        for res in res_period.filter(
            statut_reser__in=['Confirmée', 'En cours', 'Terminée'],
        ):
            overlap_start = max(res.date_arrivee, d_debut)
            overlap_end = min(res.date_depart, d_fin)
            room_nights += (overlap_end - overlap_start).days + 1
        taux_occupation = round(room_nights / capacity * 100, 1)

    depenses_par_categorie = list(
        exp_period.exclude(categorie='Toute').values('categorie').annotate(
            total=Sum('montant')
        ).order_by('-total')
    )

    reservations_par_statut = list(
        res_period.exclude(statut_reser='Toute').values('statut_reser').annotate(
            nb=Count('id')
        ).order_by('statut_reser')
    )

    nb_paiements = pay_period.count()
    montant_moyen_paiement = pay_period.aggregate(avg=Avg('montant_paye'))['avg'] or 0
    revenu_moyen_par_reservation = (recettes / nb_reservations) if nb_reservations else 0
    revenu_moyen_par_client = (recettes / nb_clients) if nb_clients else 0
    depense_moyenne_journaliere = (depenses_total / nb_jours) if nb_jours else 0
    marge_beneficiaire = round((benefice_net / recettes) * 100, 2) if recettes else 0
    revpar = round((recettes / capacity), 2) if capacity else 0
    encours_impayes = sum(float(r.solde_du or 0) for r in res_period.select_related('client', 'chambre'))

    paiements_par_methode = list(
        pay_period.values('methode_paiement').annotate(
            nb=Count('id'),
            total=Sum('montant_paye')
        ).order_by('-total')
    )

    top_clients = list(
        pay_period.values(
            'reservation__client',
            'reservation__client__nom',
            'reservation__client__prenom'
        ).annotate(
            total=Sum('montant_paye'),
            nb_paiements=Count('id')
        ).order_by('-total')[:5]
    )

    top_chambres = list(
        res_period.values(
            'chambre__numero_chambre',
            'chambre__type_chambre'
        ).annotate(
            nb_reservations=Count('id')
        ).order_by('-nb_reservations')[:5]
    )

    top_services = list(
        Service.objects.filter(
            reservation__chambre__hotel=hotel,
            reservation__date_arrivee__lte=d_fin,
            reservation__date_depart__gte=d_debut
        ).values('nom_service').annotate(
            quantite=Count('id'),
            total=Sum('prix_service')
        ).order_by('-total')[:5]
    ) if hotel else list(
        Service.objects.filter(
            reservation__date_arrivee__lte=d_fin,
            reservation__date_depart__gte=d_debut
        ).values('nom_service').annotate(
            quantite=Count('id'),
            total=Sum('prix_service')
        ).order_by('-total')[:5]
    )

    daily_finance = []
    for i in range(nb_jours):
        day = d_debut + timedelta(days=i)
        day_recettes = pay_period.filter(date_paiement__date=day).aggregate(total=Sum('montant_paye'))['total'] or 0
        day_depenses = exp_period.filter(date_depense=day).aggregate(total=Sum('montant'))['total'] or 0
        daily_finance.append({
            'date': day,
            'recettes': day_recettes,
            'depenses': day_depenses,
            'resultat': day_recettes - day_depenses,
        })

    return {
        'date_debut': date_debut,
        'date_fin': date_fin,
        'taux_occupation': taux_occupation,
        'total_recettes': recettes,
        'total_depenses': depenses_total,
        'benefice_net': benefice_net,
        'nb_reservations': nb_reservations,
        'nb_clients': nb_clients,
        'nb_paiements': nb_paiements,
        'montant_moyen_paiement': montant_moyen_paiement,
        'revenu_moyen_par_reservation': revenu_moyen_par_reservation,
        'revenu_moyen_par_client': revenu_moyen_par_client,
        'depense_moyenne_journaliere': depense_moyenne_journaliere,
        'marge_beneficiaire': marge_beneficiaire,
        'room_nights': room_nights,
        'capacity': capacity,
        'revpar': revpar,
        'encours_impayes': encours_impayes,
        'depenses_par_categorie': depenses_par_categorie,
        'reservations_par_statut': reservations_par_statut,
        'paiements_par_methode': paiements_par_methode,
        'top_clients': top_clients,
        'top_chambres': top_chambres,
        'top_services': top_services,
        'daily_finance': daily_finance,
    }


@_admin_required
def report_export_pdf(request):
    from django.template.loader import render_to_string
    from django.http import HttpResponse
    from xhtml2pdf import pisa
    from io import BytesIO

    data = _get_report_data(request)
    if not data:
        messages.error(request, "Veuillez sélectionner une période valide.")
        return redirect('reports')

    html = render_to_string('core/reports/rapport_pdf.html', data)
    result = BytesIO()
    pdf = pisa.pisaDocument(BytesIO(html.encode("UTF-8")), result)
    if not pdf.err:
        response = HttpResponse(result.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'inline; filename="rapport_{data["date_debut"]}_{data["date_fin"]}.pdf"'
        return response
    return HttpResponse("Erreur lors de la génération du PDF", status=500)


@_admin_required
def report_export_excel(request):
    from django.http import HttpResponse
    from io import BytesIO
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    data = _get_report_data(request)
    if not data:
        messages.error(request, "Veuillez sélectionner une période valide.")
        return redirect('reports')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rapport"

    # Styles
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")
    subheader_font = Font(bold=True, color="4338CA", size=11)
    subheader_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
    border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB'),
    )

    # Titre
    ws.merge_cells('A1:B1')
    ws['A1'] = f"Rapport Financier — {data['date_debut']} au {data['date_fin']}"
    ws['A1'].font = Font(bold=True, size=14, color="6366F1")

    # KPIs
    row = 3
    ws.merge_cells(f'A{row}:B{row}')
    ws[f'A{row}'] = "INDICATEURS CLES"
    ws[f'A{row}'].font = header_font
    ws[f'A{row}'].fill = header_fill

    kpis = [
        ("Taux d'occupation", f"{data['taux_occupation']}%"),
        ("Recettes", f"{data['total_recettes']:,.0f} FCFA".replace(',', '.')),
        ("Dépenses", f"{data['total_depenses']:,.0f} FCFA".replace(',', '.')),
        ("Bénéfice net", f"{data['benefice_net']:,.0f} FCFA".replace(',', '.')),
        ("Réservations", str(data['nb_reservations'])),
        ("Clients uniques", str(data['nb_clients'])),
        ("Nombre de paiements", str(data['nb_paiements'])),
        ("Paiement moyen", f"{data['montant_moyen_paiement']:,.0f} FCFA".replace(',', '.')),
        ("Revenu moyen / réservation", f"{data['revenu_moyen_par_reservation']:,.0f} FCFA".replace(',', '.')),
        ("Revenu moyen / client", f"{data['revenu_moyen_par_client']:,.0f} FCFA".replace(',', '.')),
        ("Marge bénéficiaire", f"{data['marge_beneficiaire']:.2f}%"),
        ("Nuitées vendues", str(data['room_nights'])),
        ("Capacité (nuitées)", str(data['capacity'])),
        ("RevPAR", f"{data['revpar']:,.2f} FCFA".replace(',', '.')),
        ("Encours impayés", f"{data['encours_impayes']:,.0f} FCFA".replace(',', '.')),
    ]

    for label, value in kpis:
        row += 1
        ws[f'A{row}'] = label
        ws[f'B{row}'] = value
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'A{row}'].border = border
        ws[f'B{row}'].border = border
        ws[f'B{row}'].alignment = Alignment(horizontal='right')

    # Dépenses par catégorie
    if data['depenses_par_categorie']:
        row += 2
        ws.merge_cells(f'A{row}:B{row}')
        ws[f'A{row}'] = "DEPENSES PAR CATEGORIE"
        ws[f'A{row}'].font = header_font
        ws[f'A{row}'].fill = header_fill

        row += 1
        ws[f'A{row}'] = "Catégorie"
        ws[f'B{row}'] = "Montant (FCFA)"
        ws[f'A{row}'].font = subheader_font
        ws[f'A{row}'].fill = subheader_fill
        ws[f'B{row}'].font = subheader_font
        ws[f'B{row}'].fill = subheader_fill
        ws[f'B{row}'].alignment = Alignment(horizontal='right')

        for d in data['depenses_par_categorie']:
            row += 1
            ws[f'A{row}'] = d['categorie']
            ws[f'B{row}'] = float(d['total'])
            ws[f'B{row}'].number_format = '#,##0'
            ws[f'A{row}'].border = border
            ws[f'B{row}'].border = border
            ws[f'B{row}'].alignment = Alignment(horizontal='right')

    # Réservations par statut
    if data['reservations_par_statut']:
        row += 2
        ws.merge_cells(f'A{row}:B{row}')
        ws[f'A{row}'] = "RESERVATIONS PAR STATUT"
        ws[f'A{row}'].font = header_font
        ws[f'A{row}'].fill = header_fill

        row += 1
        ws[f'A{row}'] = "Statut"
        ws[f'B{row}'] = "Nombre"
        ws[f'A{row}'].font = subheader_font
        ws[f'A{row}'].fill = subheader_fill
        ws[f'B{row}'].font = subheader_font
        ws[f'B{row}'].fill = subheader_fill
        ws[f'B{row}'].alignment = Alignment(horizontal='right')

        for r in data['reservations_par_statut']:
            row += 1
            ws[f'A{row}'] = r['statut_reser']
            ws[f'B{row}'] = r['nb']
            ws[f'A{row}'].border = border
            ws[f'B{row}'].border = border
            ws[f'B{row}'].alignment = Alignment(horizontal='right')

    # Paiements par méthode
    if data['paiements_par_methode']:
        row += 2
        ws.merge_cells(f'A{row}:C{row}')
        ws[f'A{row}'] = "PAIEMENTS PAR METHODE"
        ws[f'A{row}'].font = header_font
        ws[f'A{row}'].fill = header_fill

        row += 1
        ws[f'A{row}'] = "Méthode"
        ws[f'B{row}'] = "Transactions"
        ws[f'C{row}'] = "Montant (FCFA)"
        for col in ('A', 'B', 'C'):
            ws[f'{col}{row}'].font = subheader_font
            ws[f'{col}{row}'].fill = subheader_fill
            ws[f'{col}{row}'].border = border

        for p in data['paiements_par_methode']:
            row += 1
            ws[f'A{row}'] = p['methode_paiement']
            ws[f'B{row}'] = p['nb']
            ws[f'C{row}'] = float(p['total'] or 0)
            ws[f'A{row}'].border = border
            ws[f'B{row}'].border = border
            ws[f'C{row}'].border = border
            ws[f'B{row}'].alignment = Alignment(horizontal='right')
            ws[f'C{row}'].alignment = Alignment(horizontal='right')
            ws[f'C{row}'].number_format = '#,##0'

    # Top clients
    if data['top_clients']:
        row += 2
        ws.merge_cells(f'A{row}:C{row}')
        ws[f'A{row}'] = "TOP 5 CLIENTS"
        ws[f'A{row}'].font = header_font
        ws[f'A{row}'].fill = header_fill

        row += 1
        ws[f'A{row}'] = "Client"
        ws[f'B{row}'] = "Paiements"
        ws[f'C{row}'] = "Montant (FCFA)"
        for col in ('A', 'B', 'C'):
            ws[f'{col}{row}'].font = subheader_font
            ws[f'{col}{row}'].fill = subheader_fill
            ws[f'{col}{row}'].border = border

        for c in data['top_clients']:
            row += 1
            ws[f'A{row}'] = f"{c.get('reservation__client__prenom', '')} {c.get('reservation__client__nom', '')}".strip()
            ws[f'B{row}'] = c['nb_paiements']
            ws[f'C{row}'] = float(c['total'] or 0)
            ws[f'A{row}'].border = border
            ws[f'B{row}'].border = border
            ws[f'C{row}'].border = border
            ws[f'B{row}'].alignment = Alignment(horizontal='right')
            ws[f'C{row}'].alignment = Alignment(horizontal='right')
            ws[f'C{row}'].number_format = '#,##0'

    # Ajuster la largeur des colonnes
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 25

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="rapport_{data["date_debut"]}_{data["date_fin"]}.xlsx"'
    return response
